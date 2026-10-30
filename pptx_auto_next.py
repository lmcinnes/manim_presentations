#!/usr/bin/env python3
"""
pptx_auto_next.py - honour manim-slides' auto_next in a converted PowerPoint deck.

`manim-slides convert --to pptx` ignores `next_slide(auto_next=True)`, so
every slide waits for a click. This gives each auto_next slide a timed
advance (the length of its video), so e.g. one scene's closing fade-out runs
straight into the next scene's opening. A click still advances early.

The deck must have been converted from exactly these slide JSON files, in
this order (one pptx slide per manim-slides slide). The deck is edited in
place; run it before shrink_pptx.py (which keeps the slide XML untouched).

Example:
  pptx_auto_next.py talk.pptx slides/Intro.json slides/Middle.json slides/End.json
"""

import argparse
import json
import sys
from pathlib import Path

import av
from lxml import etree
from pptx import Presentation

P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
PAD_MS = 50  # let the last frames show before advancing


def video_ms(path: Path) -> int:
    with av.open(str(path)) as container:
        if container.duration is not None:  # in microseconds
            return round(container.duration / 1000)
        stream = container.streams.video[0]
        return round(float(stream.duration * stream.time_base) * 1000)


def resolve(file: str, json_path: Path) -> Path:
    """Slide video paths are relative to where manim-slides render ran,
    normally the parent of the slides folder."""
    for base in (Path.cwd(), json_path.parent.parent, json_path.parent):
        if (base / file).is_file():
            return base / file
    sys.exit(f"{json_path}: cannot find video {file}")


def set_auto_advance(slide, ms: int) -> None:
    sld = slide.element
    for old in sld.findall(f"{P_NS}transition"):
        sld.remove(old)
    transition = etree.Element(f"{P_NS}transition", advTm=str(ms))
    # schema order: cSld, clrMapOvr, transition, timing, extLst
    anchor = sld.find(f"{P_NS}timing")
    if anchor is None:
        anchor = sld.find(f"{P_NS}extLst")
    if anchor is not None:
        anchor.addprevious(transition)
    else:
        sld.append(transition)


def main():
    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    p.add_argument("pptx", type=Path, help="deck to edit in place")
    p.add_argument("slides", type=Path, nargs="+", help="slide JSON files, in order")
    args = p.parse_args()

    configs = [
        (cfg, path)
        for path in args.slides
        for cfg in json.loads(path.read_text())["slides"]
    ]
    prs = Presentation(str(args.pptx))
    if len(prs.slides) != len(configs):
        sys.exit(
            f"{args.pptx} has {len(prs.slides)} slides but the JSON files "
            f"describe {len(configs)}: was it converted from these files?"
        )

    changed = []
    for number, (slide, (cfg, path)) in enumerate(zip(prs.slides, configs), 1):
        if not cfg.get("auto_next") or cfg.get("loop"):
            continue
        ms = video_ms(resolve(cfg["file"], path)) + PAD_MS
        set_auto_advance(slide, ms)
        changed.append(f"{number} ({path.stem}, {ms / 1000:.2f}s)")

    prs.save(str(args.pptx))
    print(f"auto-advance on {len(changed)} slide(s)" + (":" if changed else ""))
    for line in changed:
        print("  " + line)


if __name__ == "__main__":
    main()
